import openpyxl
import os

DL = "C:/Users/Maulana/Downloads"

# Read BV reference
print("=== BV_Pembangunan_Hotel.xlsx ===")
wb = openpyxl.load_workbook(os.path.join(DL, "BV_Pembangunan_Hotel.xlsx"), data_only=True)
print("Sheets:", wb.sheetnames)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols) ---")
    for row_idx in range(1, min(ws.max_row + 1, 30)):
        row_vals = []
        for col_idx in range(1, min(ws.max_column + 1, 12)):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_vals.append(f"{cell.coordinate}={repr(val)[:50]}")
        if row_vals:
            print(f"Row {row_idx}: {row_vals}")

print("\n\n=== BV_Pembangunan_Hotel (1).xlsx ===")
wb2 = openpyxl.load_workbook(os.path.join(DL, "BV_Pembangunan_Hotel (1).xlsx"), data_only=True)
print("Sheets:", wb2.sheetnames)
for sheet_name in wb2.sheetnames:
    ws = wb2[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols) ---")
    for row_idx in range(1, min(ws.max_row + 1, 35)):
        row_vals = []
        for col_idx in range(1, min(ws.max_column + 1, 12)):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_vals.append(f"{cell.coordinate}={repr(val)[:50]}")
        if row_vals:
            print(f"Row {row_idx}: {row_vals}")

print("\n\n=== RAB_dsfsfdsf.xlsx ===")
wb3 = openpyxl.load_workbook(os.path.join(DL, "RAB_dsfsfdsf.xlsx"), data_only=True)
print("Sheets:", wb3.sheetnames)
for sheet_name in wb3.sheetnames:
    ws = wb3[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols) ---")
    for row_idx in range(1, min(ws.max_row + 1, 40)):
        row_vals = []
        for col_idx in range(1, min(ws.max_column + 1, 12)):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_vals.append(f"{cell.coordinate}={repr(val)[:50]}")
        if row_vals:
            print(f"Row {row_idx}: {row_vals}")

# Also check merged cells
print("\n\n=== Merged cells in BV (1) ===")
for sheet_name in wb2.sheetnames:
    ws = wb2[sheet_name]
    if ws.merged_cells.ranges:
        print(f"Sheet '{sheet_name}' merged: {[str(r) for r in ws.merged_cells.ranges]}")

print("\n=== Merged cells in RAB ===")
for sheet_name in wb3.sheetnames:
    ws = wb3[sheet_name]
    if ws.merged_cells.ranges:
        print(f"Sheet '{sheet_name}' merged: {[str(r) for r in ws.merged_cells.ranges]}")
