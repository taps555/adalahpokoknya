import openpyxl
import os

path = "D:/projekbaru/backend/adalahpokoknya/test_export_new.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

# Check the DIBULATKAN cell and numFmt for BQ General
ws = wb["BQ General"]
print("=== BQ General footer detail ===")
for r in [13, 14, 16, 18]:
    for c in [2, 3]:
        cell = ws.cell(row=r, column=c)
        print(f"  {cell.coordinate}: value={repr(cell.value)} numFmt={cell.number_format}")

# Check C10 (data row) numFmt
print("\n=== BQ General data row numFmt ===")
for r in [10, 11, 12]:
    cell = ws.cell(row=r, column=3)
    print(f"  {cell.coordinate}: value={repr(cell.value)} numFmt={cell.number_format}")

# Also check BQ Sipil & Interior
for sn in ["BQ Sipil", "BQ Interior"]:
    ws = wb[sn]
    print(f"\n=== {sn} footer detail ===")
    for r in [11, 12, 14, 16]:
        for c in [2, 3]:
            cell = ws.cell(row=r, column=c)
            print(f"  {cell.coordinate}: value={repr(cell.value)} numFmt={cell.number_format}")
