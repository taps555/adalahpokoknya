import openpyxl
import os

DL = "D:/projekbaru/backend/adalahpokoknya"
path = os.path.join(DL, "test_export_new.xlsx")

print("Loading:", path)
wb = openpyxl.load_workbook(path, data_only=True)
print("Sheets:", wb.sheetnames)
print("Sheet count:", len(wb.sheetnames))
