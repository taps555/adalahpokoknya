import openpyxl
import os

DL = "D:/projekbaru/backend/adalahpokoknya"
path = os.path.join(DL, "test_export_new.xlsx")

wb = openpyxl.load_workbook(path, data_only=True)
print("Sheets:", wb.sheetnames)
print()

for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"=== {sn} ({ws.max_row}r x {ws.max_column}c) merges={[str(m) for m in ws.merged_cells.ranges]} ===")
    for r in range(1, min(ws.max_row + 1, 30)):
        vals = []
        for c in range(1, min(ws.max_column + 1, 13)):
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if v is not None:
                vals.append(f"{cell.coordinate}={repr(v)[:45]}")
        if vals:
            print(f"  R{r}: {vals}")
    print()
